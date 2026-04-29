from pathlib import Path
from textwrap import dedent
from uuid import uuid4

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


WORKSPACE = Path(r"E:\E_Downloads\secure_auth_v2 2\secure_auth_v2")
OUTPUT_DOC = WORKSPACE / "SecureAuth_Project_Report_Formatted_Final_v2.docx"
FLOWCHART_TOP = WORKSPACE / "report_assets" / "image2.png"
FLOWCHART_BOTTOM = WORKSPACE / "report_assets" / "image3.png"
FLOWCHART_COMBINED = WORKSPACE / "report_assets" / "flowchart_combined.png"
LPU_LOGO = WORKSPACE / "report_assets" / "image1.png"
CODE_IMAGE_DIR = WORKSPACE / "report_assets" / "code_images"


def crop_to_content(img: Image.Image) -> Image.Image:
    rgb = img.convert("RGB")
    width, height = rgb.size
    left, top, right, bottom = width, height, 0, 0
    found = False
    for y in range(height):
        for x in range(width):
            r, g, b = rgb.getpixel((x, y))
            if (r, g, b) != (255, 255, 255):
                found = True
                left = min(left, x)
                top = min(top, y)
                right = max(right, x)
                bottom = max(bottom, y)
    if not found:
        return img
    padding = 12
    left = max(0, left - padding)
    top = max(0, top - padding)
    right = min(width, right + padding)
    bottom = min(height, bottom + padding)
    return img.crop((left, top, right, bottom))


def build_flowchart() -> Path:
    top = crop_to_content(Image.open(FLOWCHART_TOP))
    bottom = crop_to_content(Image.open(FLOWCHART_BOTTOM))
    canvas_width = max(top.width, bottom.width)
    gap = 10
    canvas = Image.new("RGB", (canvas_width, top.height + gap + bottom.height), "white")
    top_x = (canvas_width - top.width) // 2
    bottom_x = (canvas_width - bottom.width) // 2
    canvas.paste(top.convert("RGB"), (top_x, 0))
    canvas.paste(bottom.convert("RGB"), (bottom_x, top.height + gap))
    FLOWCHART_COMBINED.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(FLOWCHART_COMBINED)
    return FLOWCHART_COMBINED


def get_code_font(size=16):
    candidates = [
        r"C:\Windows\Fonts\consola.ttf",
        r"C:\Windows\Fonts\cour.ttf",
        r"C:\Windows\Fonts\lucon.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def wrap_code_line(line: str, max_chars: int):
    if len(line) <= max_chars:
        return [line]
    parts = []
    remaining = line
    while len(remaining) > max_chars:
        parts.append(remaining[:max_chars] + " \\")
        remaining = remaining[max_chars:]
    parts.append(remaining)
    return parts


def render_code_file_images(source_path: Path, out_dir: Path, max_lines=42, max_chars=92):
    out_dir.mkdir(parents=True, exist_ok=True)
    font = get_code_font(16)
    title_font = get_code_font(18)
    bg = "#f7f8fa"
    header_bg = "#1f2937"
    code_color = "#111827"
    line_no_color = "#6b7280"
    border = "#d1d5db"

    raw_lines = source_path.read_text(encoding="utf-8").splitlines()
    cleaned_lines = []
    for raw in raw_lines:
        stripped = raw.strip()
        # Omit decorative section-banner comments like "# ===== LOGIN =====" from report images
        if stripped.startswith("#") and stripped.count("=") >= 8:
            continue
        cleaned_lines.append(raw)
    raw_lines = cleaned_lines
    display_lines = []
    for idx, raw in enumerate(raw_lines, start=1):
        expanded = raw.expandtabs(4)
        wrapped = wrap_code_line(expanded, max_chars)
        for part_index, part in enumerate(wrapped):
            display_lines.append((str(idx) if part_index == 0 else "", part))

    dummy = Image.new("RGB", (10, 10), "white")
    draw = ImageDraw.Draw(dummy)
    line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + 8
    title_height = draw.textbbox((0, 0), source_path.name, font=title_font)[3] + 22
    line_no_width = draw.textbbox((0, 0), f"{len(raw_lines):>4}", font=font)[2] + 20
    code_width = max(draw.textbbox((0, 0), text or " ", font=font)[2] for _, text in display_lines) + 30
    img_width = line_no_width + code_width + 30

    images = []
    for chunk_idx in range(0, len(display_lines), max_lines):
        chunk = display_lines[chunk_idx:chunk_idx + max_lines]
        img_height = title_height + len(chunk) * line_height + 26
        img = Image.new("RGB", (img_width, img_height), bg)
        draw = ImageDraw.Draw(img)
        draw.rectangle((0, 0, img_width, title_height), fill=header_bg)
        title = source_path.name
        if len(display_lines) > max_lines:
            title += f"  ({chunk_idx // max_lines + 1})"
        draw.text((16, 10), title, font=title_font, fill="white")

        y = title_height + 10
        for line_no, code in chunk:
            draw.text((12, y), f"{line_no:>4}", font=font, fill=line_no_color)
            draw.text((line_no_width, y), code if code else " ", font=font, fill=code_color)
            y += line_height

        draw.rectangle((0, 0, img_width - 1, img_height - 1), outline=border, width=1)
        out_path = out_dir / f"{source_path.stem}_{chunk_idx // max_lines + 1}.png"
        img.save(out_path)
        images.append(out_path)
    return images


def build_code_images():
    important_files = [
        WORKSPACE / "app.py",
        WORKSPACE / "database.py",
        WORKSPACE / "security.py",
        WORKSPACE / "templates" / "base.html",
        WORKSPACE / "templates" / "login.html",
        WORKSPACE / "templates" / "register.html",
        WORKSPACE / "templates" / "otp.html",
        WORKSPACE / "templates" / "dashboard.html",
        WORKSPACE / "templates" / "sessions.html",
        WORKSPACE / "templates" / "setup_2fa.html",
        WORKSPACE / "templates" / "audit_log.html",
        WORKSPACE / "templates" / "forgot_password.html",
        WORKSPACE / "templates" / "reset_password.html",
    ]
    run_dir = CODE_IMAGE_DIR / uuid4().hex[:8]
    rendered = {}
    for path in important_files:
        rendered[path.name if path.parent.name != "templates" else f"templates/{path.name}"] = render_code_file_images(path, run_dir)
    return rendered


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_lines(cell, lines, align=WD_ALIGN_PARAGRAPH.LEFT, base_size=12):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    first = True
    for text, bold, indent in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.alignment = align
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold = bold
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(base_size)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)

    for name, size in [("Heading 1", 14), ("Heading 2", 12)]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True


def add_paragraph(doc: Document, text: str, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)


def add_numbered_point(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)


def shade_paragraph(paragraph, fill="F3F3F3"):
    p_pr = paragraph._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_code_block(doc: Document, title: str, code: str):
    add_paragraph(doc, title, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    for line in dedent(code).strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        shade_paragraph(p)
        run = p.add_run(line)
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(9)


def add_image(doc: Document, image_path: Path, width_inches=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(image_path), width=Inches(width_inches))


def add_cover_page(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Pt(12)
    p.paragraph_format.right_indent = Pt(12)
    p.paragraph_format.border_bottom = None
    r = p.add_run("Secure Authentication Module for Operating Systems")
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(20)
    r.font.color.rgb = None
    r.bold = True

    p_pr = p._element.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "4C94D8")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    doc.add_paragraph()
    add_paragraph(doc, "PROJECT REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    add_paragraph(doc, "As a Field work for Course OPERATING SYSTEMS (CSE316)", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "By", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)

    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Sr. No.", "Registration No", "Name of Students", "Roll No"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    data_rows = [
        ("1", "", "", ""),
        ("2", "", "", ""),
        ("3", "12411144", "Santlaj kumar Mehta", "52"),
    ]
    for ridx, values in enumerate(data_rows, start=1):
        for cidx, value in enumerate(values):
            set_cell_text(table.rows[ridx].cells[cidx], value, bold=(ridx == 3 and cidx > 0), align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LPU_LOGO), width=Inches(5.5))
    doc.add_paragraph()
    add_paragraph(doc, "Submitted To- Ishita", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_paragraph(doc, "Lovely Professional University", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_paragraph(doc, "Jalandhar, Punjab, India.", align=WD_ALIGN_PARAGRAPH.RIGHT)
    doc.add_page_break()


def add_toc_page(doc: Document):
    add_paragraph(doc, "Table of Contents", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_text(table.rows[0].cells[0], "Sr. No.", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(table.rows[0].cells[1], "Topic", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    rows = [
        (
            "1",
            [
                ("Project Overview", False, 0.0),
                ("1.1 Introduction", False, 0.18),
                ("1.2 Objectives of the project", False, 0.18),
                ("1.3 Scope of the project", False, 0.18),
            ],
        ),
        (
            "2",
            [
                ("Module-Wise Breakdown", False, 0.0),
                ("2.1 Authentication Core Module", False, 0.18),
                ("2.2 Security Enforcement and Recovery Module", False, 0.18),
                ("2.3 Session and User Interaction Module", False, 0.18),
            ],
        ),
        (
            "3",
            [
                ("Functionalities", False, 0.0),
                ("3.1 Authentication Core Functionalities", False, 0.18),
                ("3.2 Security Monitoring and Protection Functionalities", False, 0.18),
                ("3.3 Session, Recovery and Interface Functionalities", False, 0.18),
            ],
        ),
        (
            "4",
            [
                ("Technology Used", False, 0.0),
                ("4.1 Programming Languages", False, 0.18),
                ("4.2 Libraries and Frameworks", False, 0.18),
                ("4.3 Other Tools", False, 0.18),
            ],
        ),
        ("5", [("Revision Tracking on GitHub", False, 0.0)]),
        ("6", [("Flow Diagram", False, 0.0)]),
        (
            "7",
            [
                ("Conclusion and Future Scope", False, 0.0),
                ("7.1 Conclusion", False, 0.18),
                ("7.2 Future Scope", False, 0.18),
            ],
        ),
        ("8", [("References", False, 0.0)]),
        (
            "9",
            [
                ("Appendices", False, 0.0),
                ("A. AI-Generated Project Elaboration/Breakdown Report", False, 0.18),
                ("B. Problem Statement", False, 0.18),
                ("C. Solution Code Snippets", False, 0.18),
            ],
        ),
    ]
    for sr, topic_lines in rows:
        row = table.add_row()
        set_cell_text(row.cells[0], sr, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_lines(row.cells[1], topic_lines)
    doc.add_page_break()


def add_technology_table(doc: Document, heading: str, headers, rows):
    add_heading(doc, heading, 2)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            set_cell_text(row.cells[idx], value)
    doc.add_paragraph()


def add_main_body(doc: Document, flowchart_path: Path):
    add_heading(doc, "1. Project Overview", 1)
    add_heading(doc, "1.1 Introduction", 2)
    add_paragraph(
        doc,
        "Authentication is one of the most important jobs of an operating system. Before a user can open files, use services, or work inside the system, the operating system must confirm that the user is genuine. If this first check is weak, then the full system becomes weak. That is why secure login design is a major part of operating system security.",
    )
    add_paragraph(
        doc,
        "Many simple systems still depend only on a password. In real life, that is not always enough because passwords can be guessed, stolen, reused, or shared. This project was created to show how an operating system can follow a stronger and layered method. Instead of trusting one password only, the system checks identity in more than one step and also records what happens during the process.",
    )
    add_paragraph(
        doc,
        "The project is implemented as a Flask-based prototype, but the design idea comes from operating system security. It models important OS responsibilities such as user verification, access control, session tracking, security logging, and safe password recovery. In simple words, this project shows how a secure authentication subsystem can be organised, tested, and explained in a practical way.",
    )

    add_heading(doc, "1.2 Objectives of the Project", 2)
    objectives = [
        "To design a layered authentication workflow inspired by operating system access-control principles.",
        "To secure passwords using bcrypt hashing instead of storing them in plaintext form.",
        "To implement multi-factor authentication using time-based one-time passwords (TOTP).",
        "To reduce brute-force risk by locking accounts after repeated failed login attempts.",
        "To manage and revoke active user sessions in a controlled manner.",
        "To maintain audit logs for accountability, monitoring, and security review.",
        "To provide secure password recovery through time-limited, single-use reset tokens.",
    ]
    for item in objectives:
        add_numbered_point(doc, item)

    add_heading(doc, "1.3 Scope of the Project", 2)
    add_paragraph(
        doc,
        "The scope of this project includes the complete user authentication journey. It starts from new-user registration and continues through password validation, OTP verification, session creation, session control, password reset, and audit logging. Every major step has been included so that the system can be shown as a complete security workflow instead of a small isolated login example.",
    )
    add_paragraph(
        doc,
        "This project does not directly replace the built-in login system of a real operating system. Instead, it acts as a working model that explains how an OS-style authentication layer can be designed. Its purpose is educational and demonstrative. It helps show how security ideas from operating systems can be converted into a practical authentication module.",
    )

    add_heading(doc, "2. Module-Wise Breakdown", 1)
    add_paragraph(
        doc,
        "The Secure Authentication Module is divided into three main modules so that each part has a clear and separate responsibility. This modular structure is important in operating systems because it keeps the design organised and easier to manage. Instead of keeping all logic in one place, the system separates verification, protection, and user interaction.",
    )
    add_paragraph(
        doc,
        "This separation also improves understanding. A teacher or evaluator can clearly see which part checks identity, which part applies security rules, and which part handles user-facing actions. That is why the project was structured in three modules instead of one large combined block.",
    )

    add_heading(doc, "2.1 Authentication Core Module", 2)
    add_paragraph(
        doc,
        "Purpose: The Authentication Core Module performs the main identity checks of the system. It is the part that decides whether a user should move forward in the login process or not. This module handles username lookup, password checking, and OTP-based second-step verification.",
    )
    add_paragraph(
        doc,
        "Role in the System: In operating system terms, this module acts like the central login engine. It does not give access immediately after receiving credentials. Instead, it follows an ordered process. First it checks the user record, then it verifies the password, and then it asks for OTP. Only after these steps are successful does it allow an authenticated session to be created.",
    )
    for item in [
        "Registration of new users with username, email, hashed password, and OTP secret.",
        "Password verification using bcrypt.",
        "OTP generation and validation using time-based codes.",
        "Creation of active session records after successful authentication.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2.2 Security Enforcement and Recovery Module", 2)
    add_paragraph(
        doc,
        "Purpose: This module applies the security rules that make the system stronger. It is responsible for protecting stored credentials, limiting repeated wrong attempts, supporting safe password reset, and recording security events. In simple words, this module makes sure the login system is not only functional but also secure.",
    )
    add_paragraph(
        doc,
        "Role in the System: In an operating system, security is not only about checking who the user is. It is also about how the system behaves when something goes wrong. This module handles such situations in a controlled way. If the password is wrong many times, the account gets locked. If the user forgets the password, the reset process uses a secure token. If something important happens, it is written to the audit log for review.",
    )
    for item in [
        "Tracking failed login attempts and locking accounts after three failures.",
        "Generating cryptographically secure reset tokens with expiry control.",
        "Sending email-based OTP and password-reset notifications.",
        "Maintaining a timestamped audit log of authentication events.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2.3 Session and User Interaction Module", 2)
    add_paragraph(
        doc,
        "Purpose: This module provides the pages and actions that the user sees after or during authentication. It includes the dashboard, session view, 2FA setup, audit log view, and password-reset screens. These pages are important because they let the user interact with the security system in a controlled and understandable way.",
    )
    add_paragraph(
        doc,
        "Role in the System: After a user is authenticated, the system must still manage what happens next. The operating system must know which session is active, whether other devices are logged in, and whether the user should be allowed to close those sessions. This module supports that idea by showing session details, displaying recent security activity, and helping users configure authenticator-based OTP.",
    )
    for item in [
        "Protected dashboard shown only after successful authentication.",
        "Display and revocation of active login sessions.",
        "QR-based onboarding for authenticator applications such as Google Authenticator or Authy.",
        "Audit-log viewing for security monitoring and accountability.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Functionalities", 1)
    add_heading(doc, "3.1 Authentication Core Functionalities", 2)
    add_paragraph(
        doc,
        "The authentication core focuses on verifying identity step by step. It begins with user registration and then handles the secure login path. These features make sure that the system does not trust input blindly and that access is given only after the required checks are complete.",
    )
    add_bullet(doc, "User registration with unique username validation.")
    add_bullet(doc, "Secure password hashing using bcrypt with salting.")
    add_bullet(doc, "Password verification against stored hashes.")
    add_bullet(doc, "TOTP generation and verification as the second authentication factor.")
    add_bullet(doc, "Creation of a database-backed session only after both factors succeed.")

    add_heading(doc, "3.2 Security Monitoring and Protection Functionalities", 2)
    add_paragraph(
        doc,
        "These functions improve system safety when users make mistakes or when an attacker tries to misuse the login process. The goal is to reduce risk, maintain control, and keep evidence of important actions for later checking.",
    )
    add_bullet(doc, "Account lockout after three consecutive failed password attempts.")
    add_bullet(doc, "Single-use password reset tokens with a fifteen-minute validity window.")
    add_bullet(doc, "Anti-enumeration response during password-reset requests.")
    add_bullet(doc, "Timestamped logging of registration, login, OTP, reset, and session events.")
    add_bullet(doc, "Audit visibility that helps trace suspicious activity and review system usage.")

    add_heading(doc, "3.3 Session, Recovery, and Interface Functionalities", 2)
    add_paragraph(
        doc,
        "The final group of features deals with what the user can manage and observe. This is important because security becomes stronger when users can see active sessions, understand recent actions, and recover access safely when needed.",
    )
    add_bullet(doc, "Session listing with browser, IP address, login time, and current-session marking.")
    add_bullet(doc, "Revocation of individual sessions and logout from all active sessions.")
    add_bullet(doc, "Authenticator-app setup through QR code and manual secret key entry.")
    add_bullet(doc, "Protected dashboard that groups all security operations in one place.")
    add_bullet(doc, "Email-based recovery workflow for users who forget their password.")

    add_heading(doc, "4. Technology Used", 1)
    add_paragraph(
        doc,
        "The project was developed using a simple but effective technology stack. The selected tools make it easier to build secure login features, store user information, and demonstrate operating system security ideas in a working form. Each technology was chosen because it supports a specific part of the authentication process.",
    )
    add_technology_table(
        doc,
        "4.1 Programming Languages",
        ["Language", "Purpose"],
        [
            ("Python", "Backend logic, database coordination, authentication flow, OTP handling, and audit support"),
            ("HTML", "Structure of user-facing authentication pages"),
            ("CSS", "Styling of the login, dashboard, session, and audit interfaces"),
            ("SQL", "Persistent storage queries for users, sessions, and reset tokens"),
        ],
    )
    add_paragraph(
        doc,
        "Python was selected because it is easy to understand, quick to develop with, and well supported by security-related libraries. It also helps represent system logic clearly, which is useful in an academic project where explanation is as important as implementation.",
    )
    add_technology_table(
        doc,
        "4.2 Libraries and Frameworks",
        ["Library / Framework", "Purpose"],
        [
            ("Flask", "Web routing and application flow control"),
            ("bcrypt", "Secure password hashing and verification"),
            ("pyotp", "Time-based one-time password generation and validation"),
            ("qrcode", "QR-code generation for authenticator-app onboarding"),
            ("python-dotenv", "Loading email credentials from environment variables"),
            ("smtplib", "Sending OTP and password-reset emails"),
            ("sqlite3", "Lightweight persistent user and session storage"),
        ],
    )
    add_technology_table(
        doc,
        "4.3 Other Tools",
        ["Tool", "Purpose"],
        [
            ("SQLite", "Local database storage for users, sessions, and reset tokens"),
            ("GitHub", "Version tracking and repository hosting"),
            ("Visual Studio Code", "Code development and editing"),
            ("Web Browser", "Running and testing the authentication workflow"),
        ],
    )

    add_heading(doc, "5. Revision Tracking on GitHub", 1)
    add_paragraph(
        doc,
        "Version control is important in software projects because it helps track changes over time. It also gives a safe backup of the work and makes it easier to review how the project developed from one stage to another. In this project, GitHub was used to keep the source code organised and to preserve the revision history of the authentication system.",
    )
    add_paragraph(
        doc,
        "The repository includes the main application file, database module, security module, templates, and setup files. As features such as OTP verification, session handling, and logging were added, GitHub helped maintain a clear development path. This is useful not only for coding, but also for project reporting and presentation.",
    )
    add_paragraph(doc, "Repository Name: Secure-Authentication-Module-for-Operating-Systems", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph(doc, "GitHub Link: https://github.com/dhawan2006/Secure-Authentication-Module-for-Operating-Systems.git", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph(doc, "YouTube Video Link: To be added after final recording", align=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph(doc, "LinkedIn Video Link: To be added after final posting", align=WD_ALIGN_PARAGRAPH.LEFT)

    add_heading(doc, "6. Flow Diagram", 1)
    add_paragraph(
        doc,
        "The following flow diagram shows the full path followed by a user in the system. It begins when the application is opened and the user chooses either signup or login. From there, the flow moves through input checking, password validation, OTP generation, OTP verification, and finally dashboard access after successful authentication.",
    )
    add_paragraph(
        doc,
        "This diagram is important because it explains the project in one visual structure. It shows that authentication in this system is not a single step. It is a chain of checks and decisions. The same flowchart idea used in the earlier OS project report has been kept here so that the logic remains consistent.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(flowchart_path), width=Inches(4.9))
    add_paragraph(doc, "Figure 1. Secure Authentication workflow used in the project.", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

    add_heading(doc, "7. Conclusion and Future Scope", 1)
    add_heading(doc, "7.1 Conclusion", 2)
    add_paragraph(
        doc,
        "The Secure Authentication Module for Operating Systems successfully demonstrates a layered approach to login security. The project shows that a secure system should not depend only on one password. Instead, it should combine identity checks, second-factor verification, session control, recovery support, and logging in one complete process.",
    )
    add_paragraph(
        doc,
        "The main achievement of this project is that it goes beyond a basic login page. It presents authentication as a full operating system function. The system verifies users carefully, protects access, records events, and gives control over active sessions. In this way, the project works as a useful and practical model of OS-inspired security design.",
    )

    add_heading(doc, "7.2 Future Scope", 2)
    add_bullet(doc, "Integration of biometric or WebAuthn-based authentication.")
    add_bullet(doc, "Role-based access control for multiple categories of users.")
    add_bullet(doc, "Production-grade session hardening with HTTPS-only deployment and stronger cookie policies.")
    add_bullet(doc, "Automated anomaly detection on audit logs to identify suspicious behaviour.")
    add_bullet(doc, "Adaptation of the architecture into a deeper OS-level or PAM-style integration model.")

    add_heading(doc, "8. References", 1)
    references = [
        "Python Software Foundation. Python 3 Documentation. https://docs.python.org/3/",
        "Flask Documentation. https://flask.palletsprojects.com/",
        "bcrypt Documentation and USENIX references for password hashing.",
        "RFC 6238: TOTP - Time-Based One-Time Password Algorithm.",
        "pyotp Documentation. https://pyauth.github.io/pyotp/",
        "OWASP Authentication Cheat Sheet. https://owasp.org/",
        "SQLite Documentation. https://www.sqlite.org/docs.html",
    ]
    for ref in references:
        add_bullet(doc, ref)

    add_heading(doc, "9. Appendices", 1)
    add_heading(doc, "Appendix A: AI-Generated Project Elaboration/Breakdown Report", 1)
    add_heading(doc, "A.1 What Is This Project?", 2)
    add_paragraph(
        doc,
        "SecureAuth is a Flask-based multi-factor authentication application created as an Operating Systems project prototype. It simulates a secure login subsystem by combining identity verification, time-based OTP validation, session tracking, password-reset recovery, and activity logging.",
    )
    add_heading(doc, "A.2 Project File Structure", 2)
    add_paragraph(doc, "secure_auth_v2/", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
    for line in [
        "app.py - main routing and application control",
        "database.py - SQLite data-access layer",
        "security.py - hashing, OTP, email, QR, and logging utilities",
        "templates/ - login, dashboard, audit, session, and reset interfaces",
        "logs/auth.log - persistent authentication event log",
        "users.db - user, session, and reset-token database",
    ]:
        add_bullet(doc, line)
    add_heading(doc, "A.3 Security Highlights", 2)
    for item in [
        "Passwords are stored as bcrypt hashes rather than plaintext.",
        "OTP is used as a second factor before access is granted.",
        "Accounts are locked after repeated failed attempts.",
        "Session records can be reviewed and revoked.",
        "Password reset is token-based and time-limited.",
        "All important events are written to an audit log.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Appendix B: Problem Statement", 1)
    add_paragraph(
        doc,
        "Traditional operating system login mechanisms that rely only on static passwords are vulnerable to brute-force attacks, password theft, session abuse, and weak recovery mechanisms. The objective of this project is to design and demonstrate a secure authentication module that strengthens user verification through password hashing, multi-factor authentication, controlled session management, account lockout, secure password reset, and audit logging, thereby improving overall access-control security in an OS-inspired environment.",
    )

    add_heading(doc, "Appendix C: Solution Code Snippets", 1)
    add_paragraph(
        doc,
        "This appendix contains image-based views of the important source code used in the project. The files shown below are taken from the main application, database, security, and key template files. Sensitive environment data such as .env contents has been excluded.",
    )
    code_images = build_code_images()
    ordered_sections = [
        ("C.1 Main Application Code", ["app.py"]),
        ("C.2 Database and Security Modules", ["database.py", "security.py"]),
        (
            "C.3 Key User Interface Templates",
            [
                "templates/base.html",
                "templates/login.html",
                "templates/register.html",
                "templates/otp.html",
                "templates/dashboard.html",
                "templates/sessions.html",
                "templates/setup_2fa.html",
                "templates/audit_log.html",
                "templates/forgot_password.html",
                "templates/reset_password.html",
            ],
        ),
    ]
    for heading, file_keys in ordered_sections:
        add_heading(doc, heading, 2)
        for file_key in file_keys:
            add_paragraph(doc, file_key, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True)
            for image_path in code_images[file_key]:
                add_image(doc, image_path)


def main():
    flowchart_path = build_flowchart()
    doc = Document()
    style_document(doc)
    add_cover_page(doc)
    add_toc_page(doc)
    add_main_body(doc, flowchart_path)
    doc.save(OUTPUT_DOC)
    print(OUTPUT_DOC)


if __name__ == "__main__":
    main()
